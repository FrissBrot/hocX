"""Datenschutz-Audit PRIV-006 (Umsetzungsplan Paket 5): Datenauskunft je Person als
passwortgeschütztes ZIP mit Excel und JSON - nur Daten der Person, keine Zugangsdaten."""
from __future__ import annotations

import io
import json
import secrets
from datetime import date
from decimal import Decimal

import pytest
import pyzipper
from fastapi import HTTPException
from openpyxl import load_workbook

from app.api.routes.participants import export_participant_personal_data
from app.api.routes.users import export_my_personal_data
from app.schemas.data_export import ProtectedExportRequest
from app.services.data_export import personal
from tests.factories import (
    make_app_user,
    make_current_user,
    make_finance_account,
    make_fine,
    make_protocol,
    make_protocol_element,
    make_protocol_element_block,
    make_protocol_text,
    make_template,
    make_tenant,
)


def _setup(db):
    tenant = make_tenant(db, "Auskunft")
    anna = make_app_user(db, email="anna@example.com", first_name="Anna", last_name="Auskunft", tenant_id=tenant.id)
    ben = make_app_user(db, email="ben@example.com", first_name="Ben", last_name="Anderer", tenant_id=tenant.id)
    template = make_template(db, tenant.id)
    protocol = make_protocol(db, tenant.id, template.id, protocol_number="P-7", protocol_date=date(2026, 5, 4), status="durchgeführt")
    element = make_protocol_element(db, protocol.id)
    block = make_protocol_element_block(
        db, element.id, element_type_code="attendance",
        configuration_snapshot_json={"attendance_entries": [
            {"participant_id": anna.id, "participant_name": "Anna Auskunft", "status": "excused"},
            {"participant_id": ben.id, "participant_name": "Ben Anderer", "status": "absent"},
        ]},
    )
    make_protocol_text(db, block.id, content="Anna Auskunft übernimmt das Protokoll, Ben Anderer die Kasse.")
    account = make_finance_account(db, tenant.id)
    for person, name in ((anna, "Anna Auskunft"), (ben, "Ben Anderer")):
        fine = make_fine(db, protocol_id=protocol.id, account_id=account.id, amount=Decimal("5.00"), participant_name_snapshot=name)
        fine.participant_id = person.id
    db.flush()
    return tenant, anna, ben


def test_collect_returns_only_the_persons_own_rows(db):
    _tenant, anna, _ben = _setup(db)

    data = personal.collect(db, anna)

    assert [row["status"] for row in data["attendance"]] == ["excused"]
    assert len(data["fines"]) == 1
    assert data["person"][0]["email"] == "anna@example.com"
    assert data["mentions"] and "Anna Auskunft" in data["mentions"][0]["excerpt"]
    dumped = json.dumps(data, default=str)
    assert "Ben" not in json.dumps(data["attendance"], default=str) and "password" not in dumped


def test_archive_is_encrypted_and_readable_with_the_password(db):
    _tenant, anna, _ben = _setup(db)
    password = secrets.token_urlsafe(16)

    archive, filename = personal.build_archive(db, anna, password=password, language="de")
    try:
        assert filename.startswith("Datenauskunft-Anna Auskunft-")
        with pyzipper.AESZipFile(archive.path) as zf:
            with pytest.raises(RuntimeError):
                zf.read("daten.json", pwd=b"falsches-passwort-123")
            zf.setpassword(password.encode())
            names = zf.namelist()
            data = json.loads(zf.read("daten.json"))
            sheets = load_workbook(io.BytesIO(zf.read("Auskunft.xlsx"))).sheetnames
        assert {"Auskunft.xlsx", "daten.json", "LIESMICH.txt"} <= set(names)
        assert data["tenant"] == "Auskunft"
        assert "Anwesenheit" in sheets and "Bussen" in sheets
    finally:
        archive.path.unlink(missing_ok=True)


def test_short_password_is_rejected(db):
    _tenant, anna, _ben = _setup(db)

    with pytest.raises(ValueError):
        personal.build_archive(db, anna, password="zu-kurz", language="de")


def test_only_admins_export_others_and_everyone_exports_themselves(db):
    tenant, anna, _ben = _setup(db)
    admin = make_app_user(db, email="auskunft-admin@example.com", tenant_id=tenant.id, role_code="admin")
    password = ProtectedExportRequest(password=secrets.token_urlsafe(16))

    with pytest.raises(HTTPException) as denied:
        export_participant_personal_data(anna.public_id, password, db, make_current_user(tenant.id, role="writer", user_id=admin.id))
    assert denied.value.status_code == 403

    response = export_participant_personal_data(anna.public_id, password, db, make_current_user(tenant.id, role="admin", user_id=admin.id))
    own = export_my_personal_data(password, db, make_current_user(tenant.id, role="reader", user_id=anna.id))

    assert response.media_type == own.media_type == "application/zip"


def test_person_of_another_tenant_is_not_found(db):
    _tenant, anna, _ben = _setup(db)
    other = make_tenant(db, "Fremd")

    with pytest.raises(HTTPException) as exc:
        export_participant_personal_data(anna.public_id, ProtectedExportRequest(password=secrets.token_urlsafe(16)), db, make_current_user(other.id, role="admin"))
    assert exc.value.status_code == 404
