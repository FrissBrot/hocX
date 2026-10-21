"""Datenschutz-Audit D2 (Umsetzungsplan Paket 6): Vereinsexport für die Weiterarbeit als AES-ZIP,
asynchron gebaut, 7 Tage abrufbar, ohne Zugangsdaten; Plattform-Export nur auf Wunsch mit Zugangsdaten."""
from __future__ import annotations

import dataclasses
import io
import secrets
from datetime import UTC, date, datetime, timedelta

import pytest
import pyzipper
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import select

from app.api.routes import tenant_data_export as routes
from app.core.storage import MAIN, get_object_store
from app.models import TenantDataExport
from app.services.data_export import tenant as tenant_export
from app.services.tenant_export_service import TenantExportService
from tests.factories import (
    make_app_user,
    make_current_user,
    make_event,
    make_protocol,
    make_protocol_element,
    make_protocol_element_block,
    make_template,
    make_tenant,
)


@pytest.fixture(autouse=True)
def _isolated_storage(monkeypatch, tmp_path):
    from app.core.config import settings

    monkeypatch.setattr(settings, "storage_root", str(tmp_path / "storage"))
    monkeypatch.setattr(settings, "export_root", str(tmp_path / "storage" / "exports"))
    monkeypatch.setattr(settings, "storage_work_root", str(tmp_path / "work"))


def _setup(db):
    tenant = make_tenant(db, "Exportverein")
    admin = make_app_user(db, email="export-admin@example.com", first_name="Ada", last_name="Admin", tenant_id=tenant.id, role_code="admin", password=secrets.token_urlsafe(16))
    anna = make_app_user(db, email="export-anna@example.com", first_name="Anna", last_name="Aktiv", tenant_id=tenant.id)
    make_event(db, tenant.id, title="Sommerlager", event_date=date(2026, 7, 1))
    template = make_template(db, tenant.id)
    protocol = make_protocol(db, tenant.id, template.id, protocol_number="P-1", protocol_date=date(2026, 5, 4), status="abgeschlossen")
    element = make_protocol_element(db, protocol.id)
    make_protocol_element_block(db, element.id, element_type_code="attendance", configuration_snapshot_json={"attendance_entries": [
        {"participant_id": anna.id, "participant_name": "Anna Aktiv", "status": "late"},
    ]})
    db.flush()
    user = dataclasses.replace(make_current_user(tenant.id, role="admin", user_id=admin.id), mfa_verified=True)
    return tenant, admin, anna, user


def _open(db, job: TenantDataExport, password: str) -> pyzipper.AESZipFile:
    data = b"".join(get_object_store().iter_bytes(MAIN, job.object_key))
    zf = pyzipper.AESZipFile(io.BytesIO(data))
    zf.setpassword(password.encode())
    return zf


def test_export_needs_admin_with_mfa_and_only_one_at_a_time(db):
    tenant, admin, _anna, user = _setup(db)
    password = secrets.token_urlsafe(16)
    request = routes.TenantDataExportRequest(password=password)

    with pytest.raises(HTTPException) as no_mfa:
        routes.create_tenant_data_export(request, db, dataclasses.replace(user, mfa_verified=False))
    assert no_mfa.value.status_code == 403
    with pytest.raises(HTTPException) as writer:
        routes.create_tenant_data_export(request, db, dataclasses.replace(user, current_role="writer"))
    assert writer.value.status_code == 403

    created = routes.create_tenant_data_export(request, db, user)
    assert created.status == "queued"
    job = db.get(TenantDataExport, created.id)
    assert job.password_encrypted and password not in job.password_encrypted
    with pytest.raises(HTTPException) as busy:
        routes.create_tenant_data_export(request, db, user)
    assert busy.value.status_code == 409


def test_job_builds_a_readable_encrypted_archive_without_credentials(db):
    tenant, admin, anna, user = _setup(db)
    password = secrets.token_urlsafe(16)
    job = tenant_export.request_export(db, tenant_id=tenant.id, user_id=admin.id, password=password, include_files=True, language="de")
    db.commit()

    result = tenant_export.process_pending(db)

    db.refresh(job)
    assert result["built"] == 1 and job.status == "done"
    assert job.password_encrypted is None and job.expires_at > datetime.now(UTC) + timedelta(days=6)
    zf = _open(db, job, password)
    names = set(zf.namelist())
    assert {"Mitglieder.xlsx", "Termine.xlsx", "Anwesenheit.xlsx", "Bussen.xlsx", "LIESMICH.txt"} <= names
    members = load_workbook(io.BytesIO(zf.read("Mitglieder.xlsx"))).active
    assert [cell.value for cell in members[1]][:2] == ["Vorname", "Nachname"]
    attendance = load_workbook(io.BytesIO(zf.read("Anwesenheit.xlsx"))).active
    assert attendance.cell(1, 2).value == "04.05.2026 P-1"
    assert any(row[1] == "verspätet" for row in attendance.iter_rows(min_row=2, values_only=True))
    readme = zf.read("LIESMICH.txt").decode()
    assert "P-1" in readme  # abgeschlossenes Protokoll ohne PDF wird genannt
    everything = b"".join(zf.read(name) for name in names)
    assert admin.password_hash.encode() not in everything
    with pytest.raises(RuntimeError):
        pyzipper.AESZipFile(io.BytesIO(b"".join(get_object_store().iter_bytes(MAIN, job.object_key)))).read("LIESMICH.txt", pwd=b"falsches-passwort")


def test_photos_are_included_as_original_files_in_their_event_folder(db):
    from app.models import GalleryImage, StoredFile

    tenant, admin, _anna, _user = _setup(db)
    event = make_event(db, tenant.id, title="Herbstlager", event_date=date(2026, 9, 12))
    payload = secrets.token_bytes(64)
    get_object_store().put_bytes(db, MAIN, "uploads/export-test/bild.jpg", payload, content_type="image/jpeg")
    stored = StoredFile(tenant_id=tenant.id, original_name="bild.jpg", mime_type="image/jpeg", storage_path="uploads/export-test/bild.jpg", scan_status="clean")
    db.add(stored)
    db.flush()
    db.add(GalleryImage(tenant_id=tenant.id, stored_file_id=stored.id, event_id=event.id))
    password = secrets.token_urlsafe(16)
    job = tenant_export.request_export(db, tenant_id=tenant.id, user_id=admin.id, password=password, include_files=True, language="de")
    db.commit()

    tenant_export.process_pending(db)

    db.refresh(job)
    assert _open(db, job, password).read("Fotos/12.09.2026 Herbstlager/bild.jpg") == payload


def test_download_is_tenant_scoped_and_expires(db):
    tenant, admin, _anna, user = _setup(db)
    other = make_tenant(db, "Fremdverein")
    job = tenant_export.request_export(db, tenant_id=tenant.id, user_id=admin.id, password=secrets.token_urlsafe(16), include_files=False, language="de")
    db.commit()
    tenant_export.process_pending(db)
    db.refresh(job)
    key = job.object_key

    stranger = dataclasses.replace(make_current_user(other.id, role="admin"), mfa_verified=True)
    with pytest.raises(HTTPException) as foreign:
        routes.download_tenant_data_export(job.id, None, db, stranger)
    assert foreign.value.status_code == 404
    assert routes.download_tenant_data_export(job.id, None, db, user).status_code == 200

    job.expires_at = datetime.now(UTC) - timedelta(minutes=1)
    db.commit()
    tenant_export.process_pending(db)
    db.refresh(job)
    assert job.status == "expired" and job.object_key is None
    assert db.scalar(select(TenantDataExport.object_key).where(TenantDataExport.id == job.id)) is None
    assert key


def test_platform_export_leaves_out_credentials_unless_asked(db):
    tenant, admin, _anna, _user = _setup(db)

    manifest, _ = TenantExportService().collect(db, tenant.id, "full")
    with_credentials, _ = TenantExportService().collect(db, tenant.id, "full", include_credentials=True)

    assert all(row["password_hash"] is None for row in manifest["tables"]["app_user"])
    assert manifest["tables"]["user_mfa_factor"] == [] and manifest["include_credentials"] is False
    assert any(row["password_hash"] == admin.password_hash for row in with_credentials["tables"]["app_user"])
