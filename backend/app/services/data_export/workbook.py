"""Excel-Arbeitsmappen für Exporte: eine Tabelle je Blatt, Kopfzeile fett und fixiert, Datumswerte
als echte Datumszellen (TT.MM.JJJJ), Beträge als Zahlen."""
from __future__ import annotations

import io
from collections.abc import Iterable, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


@dataclass
class Sheet:
    title: str
    headers: Sequence[str]
    rows: list[Sequence[object]] = field(default_factory=list)


def _sheet_title(value: str, used: set[str]) -> str:
    title = "".join(ch for ch in value if ch not in "[]:*?/\\")[:31] or "Tabelle"
    base, counter = title, 2
    while title in used:
        title = f"{base[:28]} {counter}"
        counter += 1
    used.add(title)
    return title


def build(sheets: Iterable[Sheet]) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    used: set[str] = set()
    for sheet in sheets:
        ws = workbook.create_sheet(_sheet_title(sheet.title, used))
        ws.append(list(sheet.headers))
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
        widths = [len(str(header)) for header in sheet.headers]
        for row in sheet.rows:
            values = []
            for value in row:
                if isinstance(value, datetime) and value.tzinfo is not None:
                    value = value.replace(tzinfo=None)
                values.append(value)
            ws.append(values)
            for index, value in enumerate(values[: len(widths)]):
                widths[index] = max(widths[index], min(len(str(value or "")), 60))
        for index, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(index)].width = width + 2
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, datetime):
                    cell.number_format = "DD.MM.YYYY HH:MM"
                elif isinstance(cell.value, date):
                    cell.number_format = "DD.MM.YYYY"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
