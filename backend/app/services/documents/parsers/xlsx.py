"""XLSX über openpyxl (seit Release 1.2.0 ohnehin im Image: XLSX-Export der Dokument-Engine):
je Tabellenblatt eine Überschrift und eine Tabelle mit den zuletzt berechneten Zellwerten.

Formeln ohne gespeicherten Wert (Datei nie in einer Tabellenkalkulation geöffnet) bleiben leer,
Diagramme, Bilder und Formatierungen werden nicht übernommen. Grosse Blätter werden gekürzt
(Warnung ``sheet_truncated``); der Parser läuft wie alle im isolierten Kindprozess."""

from __future__ import annotations

import io
from datetime import date, datetime, time

from app.services.documents.content import ParsedContent, heading, table

PARSER_VERSION = "1"
MAX_ROWS_PER_SHEET = 5000
MAX_COLUMNS = 100
MAX_SHEETS = 50


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "WAHR" if value else "FALSCH"
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="minutes") if (value.hour or value.minute) else value.date().isoformat()
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def parse_xlsx(data: bytes) -> ParsedContent:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    blocks: list[dict] = []
    warnings: list[str] = []
    try:
        for index, sheet in enumerate(workbook.worksheets):
            if index >= MAX_SHEETS:
                warnings.append("sheets_truncated")
                break
            rows: list[list[str]] = []
            for row_index, row in enumerate(sheet.iter_rows(values_only=True)):
                if row_index >= MAX_ROWS_PER_SHEET:
                    warnings.append("sheet_truncated")
                    break
                cells = [_cell(value) for value in row[:MAX_COLUMNS]]
                if len(row) > MAX_COLUMNS and "columns_truncated" not in warnings:
                    warnings.append("columns_truncated")
                rows.append(cells)
            # Leere Randzeilen/-spalten (formatierte, aber leere Zellen) abschneiden.
            while rows and not any(rows[-1]):
                rows.pop()
            width = max((max((i + 1 for i, cell in enumerate(row) if cell), default=0) for row in rows), default=0)
            rows = [row[:width] + [""] * (width - len(row[:width])) for row in rows]
            if not rows or width == 0:
                continue
            blocks.append(heading(sheet.title, 2))
            blocks.append(table(rows))
    finally:
        workbook.close()
    return ParsedContent(parser="xlsx", parser_version=PARSER_VERSION, blocks=blocks, warnings=warnings)
