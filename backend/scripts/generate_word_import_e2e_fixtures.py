"""Reproducible, synthetic Word documents for the browser import testbook.

Run from the repository: python backend/scripts/generate_word_import_e2e_fixtures.py
Requires python-docx (already a backend dependency). No real personal data.
"""
from datetime import datetime
from io import BytesIO
from pathlib import Path
import json
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document

OUTPUT = Path(__file__).resolve().parents[2] / "frontend/e2e/fixtures/word-import"
CASES = [
    ("german-date", "14.10.2023", "2023-10-14", 2023, "Herbst 2023"),
    ("leap-day", "29.02.2024", "2024-02-29", 2023, "Schalttag 2024"),
    ("cycle-end", "31.07.2024", "2024-07-31", 2023, "Zyklusende 2024"),
    ("cycle-start", "01.08.2024", "2024-08-01", 2024, "Zyklusbeginn 2024"),
    ("year-end", "31.12.2024", "2024-12-31", 2024, "Silvester 2024"),
    ("year-start", "01.01.2025", "2025-01-01", 2024, "Neujahr 2025"),
    ("no-date", None, None, None, "Datum manuell"),
]


def archive(entries):
    result = BytesIO()
    with ZipFile(result, "w", compression=ZIP_DEFLATED) as output:
        for name, data in entries:
            info = ZipInfo(name, date_time=(2020, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            output.writestr(info, data)
    return result.getvalue()


def main():
    OUTPUT.mkdir(parents=True, exist_ok=True)
    manifest = []
    for key, printed, iso, cycle, value in CASES:
        document = Document()
        document.core_properties.created = datetime(2020, 1, 1)
        document.core_properties.modified = datetime(2020, 1, 1)
        document.add_paragraph(f"Hock vom {printed}" if printed else "Hock ohne Datumsangabe")
        document.add_heading("E2E Historische Aufgaben", level=1)
        table = document.add_table(rows=1, cols=2)
        for cell, text in zip(table.rows[0].cells, ["Aufgabe", "Wert"]):
            cell.text = text
        for row in [("Feuer", value), ("Küche", f"Menü – {value}")]:
            for cell, text in zip(table.add_row().cells, row):
                cell.text = text
        raw = BytesIO()
        document.save(raw)
        with ZipFile(BytesIO(raw.getvalue())) as source:
            content = archive([(name, source.read(name)) for name in source.namelist()])
        filename = f"{key}.docx"
        (OUTPUT / filename).write_bytes(content)
        manifest.append({"file": filename, "printedDate": printed, "protocolDate": iso, "cycleYear": cycle,
                         "rows": [["Feuer", value], ["Küche", f"Menü – {value}"]]})
    # Deliberately reverse the chronology and include a non-importable attachment.
    names = ["cycle-start.docx", "cycle-end.docx", "leap-day.docx"]
    (OUTPUT / "historical-batch.zip").write_bytes(archive(
        [(name, (OUTPUT / name).read_bytes()) for name in names] + [("README.txt", b"Synthetic E2E attachment")]
    ))
    (OUTPUT / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n")


if __name__ == "__main__":
    main()
